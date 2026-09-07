#!/usr/bin/env python3
"""Build the payout comparison sheets from tools/paytables.json.

    python3 tools/build-payout-sheets.py

Outputs (both regenerated from scratch every run):
    video-poker-payouts.html   interactive — tick variations, rows self-hide,
                               live heat map, switchable baseline
    video-poker-payouts.xlsx   every variation, Excel column outlining
                               collapsed to Standard full-pay, static heat map

Run tools/extract-paytables.js first so the JSON reflects the current game code.
openpyxl is only needed for the .xlsx; the .html is written either way.
"""
import json, math, os, sys, html

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = json.load(open(os.path.join(ROOT, 'tools', 'paytables.json'), encoding='utf-8'))

HANDS    = DATA['hands']
BASELINE = DATA['baseline']
GAMES    = DATA['games']

# flat list of every variation, in game order
COLS = []
for g in GAMES:
    for i, t in enumerate(g['tables']):
        COLS.append({
            'key': f"{g['key']}:{i}", 'game': g['name'], 'gameKey': g['key'],
            'type': g['type'], 'label': t['label'], 'ret': t['return'],
            'pays': t['pays'], 'top': i == 0,
        })

# Heat map: light green when a hand pays better than the baseline game, light
# red when worse. Each ROW gets its own scale, because one scale across the
# whole grid is useless — Four Aces swings by 1075 coins while Full House
# swings by 30, so on a shared scale every Full House cell rounds to white.
GREEN, RED, GAMMA = (188, 226, 191), (245, 195, 195), 0.5

def row_scale(hand, bcol):
    """Biggest gap from the baseline anywhere on this row.

    Measured across ALL variations, not just the ones on screen, so a cell
    keeps the same shade whatever is selected and a genuinely small gap still
    looks small.
    """
    m = 0
    for c in COLS:
        d = cell_delta(c['pays'], bcol['pays'], hand)
        if d is not None:
            m = max(m, abs(d))
    return m or None

def heat_rgb(d, scale):
    if d is None or not scale:
        return None
    if abs(d) < 1e-9:
        return None
    # sqrt curve: without it the many small gaps sit near white and read as "no
    # difference" when they are really a fifth of the row's spread
    i = min(1.0, abs(d) / scale) ** GAMMA
    tgt = GREEN if d > 0 else RED
    return tuple(round(255 + (c - 255) * i) for c in tgt)

# Rows are often absent from a game not because it pays nothing for those
# cards, but because it files them under another name or splits them across
# tiers. Every wild game pays a royal — it just calls it "Natural Royal" — and
# Bonus Poker has no "Four of a Kind" line yet certainly pays four of a kind,
# under three separate tiers. Colouring those blanks as a missing payout would
# be a lie, so hands are grouped into families and a game that pays ANY member
# counts as paying the hand. A real gap (no pair pays anything in Deuces Wild)
# has no paying family member and still colours.
#
# Most families fall straight out of the baseline map, which already sends
# every quad tier and every royal to one Jacks or Better line. The
# five-of-a-kinds have no Jacks or Better equivalent to be mapped to, so that
# family is named here.
EXTRA_FAMILIES = [['Five of a Kind', 'Five Aces', 'Five 3s–5s', 'Five 6s–Ks']]

def _families():
    by_eq = {}
    for hand in HANDS:
        by_eq.setdefault(BASELINE.get(hand) or hand, set()).add(hand)
    for extra in EXTRA_FAMILIES:
        merged = set()
        for h in extra:
            merged |= by_eq.pop(BASELINE.get(h) or h, set())
        by_eq[extra[0]] = merged
    fam = {}
    for members in by_eq.values():
        for h in members:
            fam[h] = sorted(members)
    return fam

FAMILY = _families()

def covers(pays, hand):
    """Does this game pay the hand, or a sibling row standing in for it?"""
    return any(pays.get(h) is not None for h in FAMILY.get(hand, [hand]))

def cell_delta(pays, bpays, hand):
    """How far this cell sits from the baseline, or None to leave it neutral.

    A hand with a payout is judged against the baseline's EQUIVALENT line, so
    every quad tier compares with the baseline's single "Four of a Kind" —
    that comparison is the whole point of a bonus game.

    A blank cell is only a gap when the baseline lists that same hand by name.
    Two blanks are not a difference: with a Jacks or Better baseline nothing
    pays "Four Deuces", so that whole row must stay white rather than judging
    every blank against the baseline's quad line.
    """
    v = pays.get(hand)
    if v is None:
        if covers(pays, hand):
            return None            # a sibling row stands in for it — naming only
        b = bpays.get(hand)
        return None if b is None else -b
    equiv = BASELINE.get(hand) or hand
    b = bpays.get(equiv)
    if b is None:
        if covers(bpays, equiv):
            return None
        b = 0
    return v - b

BASE_COL = next(c for c in COLS if c['gameKey'] == DATA['baselineGame'] and c['top'])

NOTES = [
    'Payouts are the MAX-BET column (5 coins) — what the machine pays for a full 5-coin wager. '
    'The Amounts button switches to per-coin figures (everything ÷ 5), which is how pay tables are '
    'usually quoted: a royal reads 800 rather than 4,000, making the ratios between hands easier to '
    'compare. Shading does not change, since dividing every value and its baseline by 5 leaves the '
    'ratios identical.',
    '"—" means the game does not offer that hand at all, which is not the same as paying zero.',
    'Rows run roughly in descending payout order, using the highest each hand reaches across all '
    'variations. Games rank hands differently, so the order is indicative, not exact for any one game.',
    'Heat map compares each payout with the baseline game\'s equivalent hand: green pays better, '
    'red pays worse. Each ROW is scaled to its own widest gap, so a 5-coin difference on Full House '
    'is as readable as a 1,075-coin one on Four Aces. That scale always spans all ' + str(len(COLS)) + ' variations, '
    'not just the ones on screen, so a shade never changes when you change the selection. Every '
    'quad tier is judged against the baseline\'s single "Four of a Kind" line, which is the point '
    'of a bonus game.',
    'A hand a game does not pay at all is shaded too: red where the baseline pays it and this game '
    'does not (no pair pays anything in Deuces), green where this game pays something the baseline '
    'never does (five of a kind). Two blanks are not a difference, so a row the baseline does not '
    'pay either stays white across the board.',
    'A blank left by naming rather than by the pay table is neutral. Hands are grouped into '
    'families — the royals, the pair minimums, the twelve quad tiers, the five-of-a-kinds — and a '
    'game paying any member counts as paying the hand. Every wild game pays a royal, it just calls '
    'it "Natural Royal"; Tens or Better pays a pair of jacks under its own name; Bonus Poker has no '
    '"Four of a Kind" line yet pays four of a kind under three tiers.',
    'Optimal return is the exact expected value under perfect play, computed over every possible '
    'deal — not a simulation. All ' + str(len(COLS)) + ' tables match their published Wizard of Odds figures.',
    'Generated by tools/build-payout-sheets.py from tools/paytables.json, which is itself read out '
    'of index.html. Nothing here is transcribed by hand. Re-run both to refresh after a rules change.',
]

# ─────────────────────────────── HTML ────────────────────────────────

def build_html(dest):
    payload = {
        'hands': HANDS, 'baseline': BASELINE, 'generated': DATA['generated'],
        'cols': [{k: c[k] for k in ('key', 'game', 'gameKey', 'type', 'label', 'ret', 'pays', 'top')}
                 for c in COLS],
        'defaultBaseline': BASE_COL['key'],
        'green': GREEN, 'red': RED, 'gamma': GAMMA, 'family': FAMILY,
        'notes': NOTES,
    }
    doc = HTML_TEMPLATE.replace('/*__DATA__*/null', json.dumps(payload, ensure_ascii=False))
    with open(dest, 'w', encoding='utf-8') as f:
        f.write(doc)
    return dest


HTML_TEMPLATE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Video Poker — Payout Comparison</title>
<style>
:root{
  --ink:#1f2937; --muted:#6b7280; --line:#e5e7eb; --head:#1f3864;
  --bg:#ffffff; --panel:#f8fafc; --accent:#1f3864;
}
*{box-sizing:border-box}
html,body{height:100%}
body{margin:0;font:14px/1.45 -apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif;
     color:var(--ink);background:var(--bg);display:flex;flex-direction:column}
header{padding:10px 14px 9px;border-bottom:1px solid var(--line);flex:0 0 auto}
.hrow{display:flex;align-items:center;gap:12px}
.hrow > div{min-width:0;flex:1}
h1{margin:0;font-size:17px;letter-spacing:.2px}
.sub{color:var(--muted);font-size:11.5px}
.panel{padding:12px 20px;background:var(--panel);border-bottom:1px solid var(--line);
       position:relative;flex:0 0 auto;overflow:auto;max-height:56vh}
.panel.hidden{display:none}
.x{position:absolute;top:8px;right:10px;width:26px;height:26px;padding:0;line-height:1;
   font-size:17px;color:var(--muted);border-radius:50%}
.x:hover{background:#e5e7eb;color:var(--ink)}
.row.presets{padding-right:34px}
.row{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin-bottom:8px}
.row:last-child{margin-bottom:0}
button{font:inherit;font-size:12px;padding:5px 11px;border:1px solid #cbd5e1;background:#fff;
       border-radius:6px;cursor:pointer;color:var(--ink)}
button:hover{background:#eef2ff;border-color:#a5b4fc}
button.on{background:var(--accent);border-color:var(--accent);color:#fff}
#backBtn{border-color:#e5b4b4;color:#9f1239}
#backBtn:hover{background:#fee2e2;border-color:#fca5a5}
label.lbl{font-size:12px;color:var(--muted);margin-right:2px}
select{font:inherit;font-size:12px;padding:4px 6px;border:1px solid #cbd5e1;border-radius:6px;background:#fff}
details{margin-top:6px}
summary{cursor:pointer;font-size:12px;color:var(--accent);font-weight:600;outline:none}
.games{display:grid;grid-template-columns:repeat(auto-fill,minmax(255px,1fr));gap:6px 18px;margin-top:9px}
.game{border:1px solid var(--line);border-radius:7px;padding:7px 9px;background:#fff}
.game h3{margin:0 0 4px;font-size:12px;display:flex;justify-content:space-between;gap:6px;align-items:baseline}
.tag{font-size:9.5px;font-weight:700;letter-spacing:.4px;text-transform:uppercase;
     padding:1px 5px;border-radius:9px;white-space:nowrap}
.t-Standard{background:#e5e7eb;color:#374151}
.t-Deuces{background:#d1fae5;color:#065f46}
.t-Joker{background:#fef3c7;color:#92400e}
.game label{display:flex;align-items:center;gap:6px;font-size:12px;padding:1px 0;cursor:pointer}
.game label span.pct{margin-left:auto;color:var(--muted);font-variant-numeric:tabular-nums;font-size:11px}
.wrap{flex:1 1 auto;overflow:auto;min-height:0}
table{border-collapse:separate;border-spacing:0;font-variant-numeric:tabular-nums}
th,td{padding:4px 9px;white-space:nowrap;border-bottom:1px solid var(--line)}
thead th{position:sticky;background:var(--head);color:#fff;text-align:center;font-size:11.5px;font-weight:600}
thead tr.r1 th{top:0;z-index:6}
thead tr.r2 th{top:26px;z-index:6;font-weight:400;font-size:10.5px;opacity:.86}
thead tr.r3 th{top:47px;z-index:6;font-weight:600;font-size:11px;background:#2c4a86}
th.hand,td.hand{position:sticky;left:0;text-align:left;font-weight:600;background:#fff;z-index:5;
                border-right:1px solid var(--line)}
thead th.hand{background:var(--head);z-index:7}
tbody tr:nth-child(even) td.hand{background:#fafafa}
tbody tr:hover td{outline:1px solid #c7d2fe;outline-offset:-1px}
td.v{text-align:right;font-size:12.5px}
td.na{text-align:center;color:#cbd5e1}
td.base{box-shadow:inset 2px 0 0 #1f3864,inset -2px 0 0 #1f3864}
.empty{padding:40px 20px;color:var(--muted);text-align:center}
.legend{display:flex;align-items:center;gap:7px;font-size:11px;color:var(--muted)}
.sw{width:17px;height:12px;border:1px solid #d1d5db;border-radius:2px;display:inline-block}
footer{padding:14px 20px 26px;color:var(--muted);font-size:11.5px;max-width:1000px}
footer li{margin-bottom:5px}
@media (max-width:640px){
  .panel{padding:10px 12px}
  th,td{padding:3px 6px}
  h1{font-size:15px}
}
@media print{.panel,header button{display:none}.wrap{max-height:none;overflow:visible}}
</style></head><body>
<header>
  <div class="hrow">
    <div>
      <h1>Video Poker — Payout Comparison</h1>
      <div class="sub" id="sub"></div>
    </div>
    <button id="optBtn" class="on">Hide options</button>
    <button id="backBtn" title="Back to the game">&times;&nbsp; Close</button>
  </div>
</header>

<div class="panel" id="panel">
  <button class="x" id="closeBtn" title="Hide options" aria-label="Hide options">&times;</button>
  <div class="row presets">
    <span class="lbl">Show:</span>
    <button data-pick="standard">Standard only</button>
    <button data-pick="top">Full-pay of each game</button>
    <button data-pick="all">All <span id="allN"></span></button>
    <button data-pick="none">None</button>
    <span style="flex:1"></span>
    <span class="lbl">Amounts</span>
    <button id="scaleBtn">Max bet (5 coins)</button>
    <span class="lbl">Heat map vs</span>
    <select id="base"></select>
    <button id="heatToggle" class="on">Heat map on</button>
  </div>
  <div class="row legend">
    <span class="sw" id="sw-lo"></span><span>pays less</span>
    <span class="sw" style="background:#fff"></span><span>same</span>
    <span class="sw" id="sw-hi"></span><span>pays more</span>
    <span style="margin-left:10px">Each row is shaded against its own widest gap, measured across all
    <span class="allN"></span> variations. Rows with nothing to show in the ticked columns hide themselves.</span>
  </div>
  <details id="picker"><summary>Choose variations individually</summary>
    <div class="games" id="games"></div>
  </details>
</div>

<!-- notes live INSIDE the scroll area: as a flex sibling their height starved
     the table, which is what you notice first on a phone -->
<div class="wrap">
  <div id="table"></div>
  <footer><strong>Notes</strong><ul id="notes"></ul></footer>
</div>

<script>
const D = /*__DATA__*/null;
const $ = s => document.querySelector(s);
let sel = new Set(), baseKey = D.defaultBaseline, heatOn = true;
// Divide the max-bet figures by the 5 coins wagered, so every number reads as
// the per-coin multiple a pay table quotes — a royal shows 800, not 4,000.
// Shading is unaffected: scaling every value and its baseline by the same
// factor leaves the ratios identical.
let perCoin = false;
const COINS = 5;
const shown = v => perCoin ? v / COINS : v;

const colOf = k => D.cols.find(c => c.key === k);
const mix = (rgb, i) => `rgb(${rgb.map(c => Math.round(255 + (c - 255) * i)).join(',')})`;

// A hand missing from a game is usually a real gap worth showing — but often
// the game only files it under another name or splits it across tiers (every
// wild game pays a royal, it just calls it "Natural Royal"; Bonus Poker has no
// "Four of a Kind" line but pays four of a kind under three tiers). A game
// paying any member of the hand's family counts as paying the hand.
const covers = (col, hand) =>
  (D.family[hand] || [hand]).some(h => col.pays[h] != null);
// How far a cell sits from the baseline, or null to leave it neutral. A hand
// WITH a payout is judged against the baseline's equivalent line, so every
// quad tier compares with the baseline's single "Four of a Kind". A BLANK is
// only a gap when the baseline lists that same hand by name — two blanks are
// not a difference, so with a Jacks or Better baseline the "Four Deuces" row
// stays white instead of judging every blank against its quad line.
function cellDelta(col, bcol, hand) {
  const v = col.pays[hand];
  if (v == null) {
    if (covers(col, hand)) return null;
    const b = bcol.pays[hand];
    return b == null ? null : -b;
  }
  const eq = D.baseline[hand] || hand;
  let b = bcol.pays[eq];
  if (b == null) {
    if (covers(bcol, eq)) return null;
    b = 0;
  }
  return v - b;
}
// Each row is shaded against its OWN widest gap — Four Aces swings by 1075
// coins and Full House by 30, so one shared scale renders every Full House
// cell white. The scale spans every variation, not just the ticked ones, so
// a shade never shifts when the selection changes.
function rowScale(hand, bcol) {
  let m = 0;
  for (const c of D.cols) {
    const d = cellDelta(c, bcol, hand);
    if (d != null) m = Math.max(m, Math.abs(d));
  }
  return m || null;
}
function heat(d, scale) {
  if (!heatOn || d == null || !scale) return null;
  if (Math.abs(d) < 1e-9) return null;
  return mix(d > 0 ? D.green : D.red, Math.pow(Math.min(1, Math.abs(d) / scale), D.gamma));
}

function pick(mode) {
  sel = new Set(
    mode === 'all'      ? D.cols.map(c => c.key)
  : mode === 'none'     ? []
  : mode === 'top'      ? D.cols.filter(c => c.top).map(c => c.key)
  : /* standard */        D.cols.filter(c => c.type === 'Standard' && c.top).map(c => c.key));
  render();
}

function buildPicker() {
  const byGame = [];
  for (const c of D.cols) {
    let g = byGame.find(x => x.key === c.gameKey);
    if (!g) byGame.push(g = { key: c.gameKey, name: c.game, type: c.type, cols: [] });
    g.cols.push(c);
  }
  $('#games').innerHTML = byGame.map(g => `
    <div class="game"><h3>${esc(g.name)}
      <span class="tag t-${g.type.split(' ')[0]}">${esc(g.type)}</span></h3>
      ${g.cols.map(c => `<label><input type="checkbox" data-k="${c.key}">
        <span>${esc(c.label)}</span>
        <span class="pct">${c.ret ? c.ret.toFixed(2) + '%' : ''}</span></label>`).join('')}
    </div>`).join('');
  $('#games').addEventListener('change', e => {
    const k = e.target.dataset.k; if (!k) return;
    e.target.checked ? sel.add(k) : sel.delete(k);
    render();
  });
  $('#base').innerHTML = D.cols.map(c =>
    `<option value="${c.key}">${esc(c.game)} — ${esc(c.label)}</option>`).join('');
  $('#base').value = baseKey;
  $('#base').addEventListener('change', e => { baseKey = e.target.value; render(); });
  $('#heatToggle').addEventListener('click', e => {
    heatOn = !heatOn;
    e.target.classList.toggle('on', heatOn);
    e.target.textContent = heatOn ? 'Heat map on' : 'Heat map off';
    render();
  });
  $('#scaleBtn').addEventListener('click', e => {
    perCoin = !perCoin;
    e.target.classList.toggle('on', perCoin);
    e.target.textContent = perCoin ? 'Per coin (÷5)' : 'Max bet (5 coins)';
    render();
  });
  document.querySelectorAll('[data-pick]').forEach(b =>
    b.addEventListener('click', () => pick(b.dataset.pick)));
  // Options panel collapses so the table gets the whole screen on a phone
  const setPanel = on => {
    $('#panel').classList.toggle('hidden', !on);
    $('#optBtn').classList.toggle('on', on);
    $('#optBtn').textContent = on ? 'Hide options' : 'Options';
  };
  $('#optBtn').addEventListener('click', () => setPanel($('#panel').classList.contains('hidden')));
  $('#closeBtn').addEventListener('click', () => setPanel(false));
  // Opened from the game's Settings in a new tab, so close it; if the browser
  // refuses (page wasn't script-opened) fall back to loading the game itself.
  $('#backBtn').addEventListener('click', () => {
    window.close();
    setTimeout(() => { location.href = './'; }, 150);
  });
  if (window.innerWidth < 640) setPanel(false);   // start out of the way on a phone
  $('#sw-lo').style.background = mix(D.red, 1);
  $('#sw-hi').style.background = mix(D.green, 1);
  $('#notes').innerHTML = D.notes.map(n => `<li>${esc(n)}</li>`).join('');
}

const esc = s => String(s).replace(/[&<>"]/g, c => ({ '&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;' }[c]));

function render() {
  document.querySelectorAll('#games input').forEach(i => { i.checked = sel.has(i.dataset.k); });
  const cols = D.cols.filter(c => sel.has(c.key));
  const bcol = colOf(baseKey);
  $('#sub').textContent =
    `${cols.length} of ${D.cols.length} variations · ${D.hands.length} hands · ` +
    `${perCoin ? 'per-coin payouts' : 'max-bet (5-coin) payouts'} · generated ${D.generated}`;

  if (!cols.length) { $('#table').innerHTML = '<div class="empty">No variations ticked.</div>'; return; }
  // a row earns its place only if some ticked column actually pays it
  const rows = D.hands.filter(h => cols.some(c => c.pays[h] != null));

  let t = '<table><thead>';
  t += '<tr class="r1"><th class="hand">Hand</th>' +
       cols.map(c => `<th>${esc(c.game)}</th>`).join('') + '</tr>';
  t += '<tr class="r2"><th class="hand">Pay table</th>' +
       cols.map(c => `<th>${esc(c.label)}</th>`).join('') + '</tr>';
  t += '<tr class="r3"><th class="hand">Optimal return</th>' +
       cols.map(c => `<th>${c.ret ? c.ret.toFixed(2) + '%' : '—'}</th>`).join('') + '</tr>';
  t += '</thead><tbody>';
  for (const h of rows) {
    const hs = rowScale(h, bcol);
    t += `<tr><td class="hand">${esc(h)}</td>`;
    for (const c of cols) {
      const v = c.pays[h];
      const bg = heat(cellDelta(c, bcol, h), hs);
      const sty = bg ? ` style="background:${bg}"` : '';
      if (v == null) { t += `<td class="na"${sty}>—</td>`; continue; }
      const isBase = c.key === baseKey;
      t += `<td class="v${isBase ? ' base' : ''}"${sty}>` + shown(v).toLocaleString() + '</td>';
    }
    t += '</tr>';
  }
  $('#table').innerHTML = t + '</tbody></table>';
}

buildPicker();
for (const el of document.querySelectorAll('#allN, .allN')) el.textContent = D.cols.length;
pick('standard');
</script></body></html>
"""

# ─────────────────────────────── XLSX ────────────────────────────────

def build_xlsx(dest):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = 'Payout Comparison'
    F = 'Arial'
    DARK = PatternFill('solid', fgColor='1F3864')
    MID  = PatternFill('solid', fgColor='2C4A86')
    TYPEC = {'Standard': 'D9D9D9', 'Deuces Wild': 'C6E0B4', 'Joker Wild': 'FFE699'}
    thin = Side(style='thin', color='D9D9D9')
    med  = Side(style='medium', color='1F3864')

    R_GAME, R_TBL, R_TYPE, R_RET, R_FIRST = 1, 2, 3, 4, 6

    for r, lbl in ((R_GAME, 'Hand'), (R_TBL, 'Pay table'), (R_TYPE, 'Game type'), (R_RET, 'Optimal return')):
        c = ws.cell(r, 1, lbl)
        c.font = Font(F, bold=True, size=(11 if r == R_GAME else 9), italic=(r != R_GAME),
                      color='FFFFFF' if r == R_GAME else '404040')
        if r == R_GAME: c.fill = DARK
        c.alignment = Alignment(horizontal='left', vertical='bottom')
    ws.cell(R_RET, 1).border = Border(bottom=med)

    # one column per variation; merge the game name across its variants
    col_of = {}
    start = 2
    for g in GAMES:
        n = len(g['tables'])
        if n > 1:
            ws.merge_cells(start_row=R_GAME, start_column=start, end_row=R_GAME, end_column=start + n - 1)
        c = ws.cell(R_GAME, start, g['name'])
        c.font = Font(F, bold=True, size=10, color='FFFFFF'); c.fill = DARK
        c.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
        for j in range(n):
            ws.cell(R_GAME, start + j).fill = DARK
        for i, t in enumerate(g['tables']):
            col = start + i
            col_of[f"{g['key']}:{i}"] = col
            c = ws.cell(R_TBL, col, t['label'])
            c.font = Font(F, size=9, color='FFFFFF'); c.fill = MID
            c.alignment = Alignment(horizontal='center', wrap_text=True)
            c = ws.cell(R_TYPE, col, g['type'])
            c.font = Font(F, size=8, bold=True)
            c.fill = PatternFill('solid', fgColor=TYPEC[g['type']])
            c.alignment = Alignment(horizontal='center', wrap_text=True)
            c = ws.cell(R_RET, col, (t['return'] or 0) / 100)
            c.font = Font(F, size=9, bold=True, color='1F3864')
            c.number_format = '0.00%'; c.alignment = Alignment(horizontal='center')
            c.border = Border(bottom=med)
            # Excel cannot filter columns, so use its outline instead: the
            # full-pay Standard tables stay open, everything else collapses.
            L = get_column_letter(col)
            wild = g['type'] != 'Standard'
            lvl = (1 if i == 0 else 2) if wild else (0 if i == 0 else 1)
            ws.column_dimensions[L].width = 12.5
            if lvl:
                ws.column_dimensions[L].outlineLevel = lvl
                ws.column_dimensions[L].hidden = True
        start += n
    LAST = start - 1

    bcol = next(c for c in COLS if c['key'] == BASE_COL['key'])
    for j, hand in enumerate(HANDS):
        r = R_FIRST + j
        c = ws.cell(r, 1, hand)
        c.font = Font(F, size=10, bold=True)
        c.border = Border(right=thin, bottom=thin)
        scale = row_scale(hand, bcol)
        for col in COLS:
            cc = ws.cell(r, col_of[col['key']])
            v = col['pays'].get(hand)
            rgb = heat_rgb(cell_delta(col['pays'], bcol['pays'], hand), scale)
            if v is None:
                cc.value = '—'
                cc.font = Font(F, size=10, color='BFBFBF')
                cc.alignment = Alignment(horizontal='center')
            else:
                cc.value = v
                cc.font = Font(F, size=10); cc.number_format = '#,##0'
                cc.alignment = Alignment(horizontal='right')
            if rgb:
                cc.fill = PatternFill('solid', fgColor='%02X%02X%02X' % rgb)
            cc.border = Border(bottom=thin)

    n = R_FIRST + len(HANDS) + 1
    head = ws.cell(n, 1, 'Notes'); head.font = Font(F, size=9, bold=True, color='1F3864')
    extra = [
        f"Baseline for the heat map: {BASE_COL['game']} {BASE_COL['label']}.",
        'Only the full-pay Standard tables are shown when the file opens. Use the outline "+" '
        'buttons above the column letters (or Data > Group and Outline > Show Detail) to reveal '
        'the other variations and the wild-card games.',
        'Excel cannot filter columns or hide rows based on a column selection — for that, open '
        'video-poker-payouts.html, which ticks variations on and off and hides empty rows itself.',
        'This sheet is a generated snapshot and contains no formulas. Re-run '
        'tools/extract-paytables.js then tools/build-payout-sheets.py to refresh it.',
    ]
    for i, txt in enumerate(NOTES + extra):
        c = ws.cell(n + 1 + i, 1, txt)
        c.font = Font(F, size=9, color='595959')

    ws.freeze_panes = ws.cell(R_FIRST, 2)
    ws.column_dimensions['A'].width = 24
    ws.row_dimensions[R_GAME].height = 30
    ws.row_dimensions[R_TBL].height = 28
    ws.row_dimensions[R_TYPE].height = 22
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryRight = True
    wb.save(dest)
    return dest


if __name__ == '__main__':
    h = build_html(os.path.join(ROOT, 'video-poker-payouts.html'))
    print(f'wrote {os.path.relpath(h, ROOT)} — {len(COLS)} variations, {len(HANDS)} hands (interactive)')
    try:
        x = build_xlsx(os.path.join(ROOT, 'video-poker-payouts.xlsx'))
        print(f'wrote {os.path.relpath(x, ROOT)} — {len(COLS)} variations, {len(HANDS)} hands')
    except ImportError:
        print('SKIPPED the .xlsx: openpyxl is not installed.\n'
              '  python3 -m venv tools/.venv && tools/.venv/bin/pip install openpyxl\n'
              '  then re-run with tools/.venv/bin/python', file=sys.stderr)
        sys.exit(2)
